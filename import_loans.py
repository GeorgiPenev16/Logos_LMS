#!/usr/bin/env python3
"""
Loan Import Script — Bulk-import historical/backdated loans into Odoo via XML-RPC.

Usage:
    python import_loans.py                  # Import loans from Excel (TEST_MODE by default)
    python import_loans.py --template       # Generate sample Excel template

Loans may have past start dates with some installments already paid.
The script bypasses the full UI workflow and directly creates records
with the correct status and journal entries.
"""

import sys
import xmlrpc.client
from datetime import date, datetime

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)

# ──────────────────────────────────────────────
# CONFIGURATION
# ──────────────────────────────────────────────

ODOO_URL = "http://localhost:8069"
ODOO_DB = "odoo_db"
ODOO_USER = "admin"
ODOO_PASS = "admin"

TEST_MODE = True  # When True, validates data but does NOT create any records

EXCEL_FILE = "loans_to_import.xlsx"

# Installments with emi_date <= CUTOFF_DATE will be marked as paid.
# Set to None to use today's date.
CUTOFF_DATE = None  # e.g. "2025-01-31" or None for today

# ──────────────────────────────────────────────
# TEMPLATE GENERATION
# ──────────────────────────────────────────────

TEMPLATE_COLUMNS = [
    "customer_egn",
    "loan_type_name",
    "app_date",
    "approval_date",
    "loan_amount",
    "term",
    "installment_type",
    "interest_rate",
    "start_date",
    "installment_start_date",
    "bank_name",
    "bank_account",
    "bank_branch_code",
    "bank_swift",
    "is_penalty",
    "penalty_type",
    "penalty_amount",
    "penalty_percentage",
    "paid_through_date",
]

EXAMPLE_ROWS = [
    {
        "customer_egn": "8501011234",
        "loan_type_name": "Потребителски кредит",
        "app_date": "2024-06-15",
        "approval_date": "2024-06-16",
        "loan_amount": 10000.00,
        "term": 12,
        "installment_type": "monthly",
        "interest_rate": 12.5,
        "start_date": "2024-07-01",
        "installment_start_date": "2024-07-01",
        "bank_name": "УниКредит Булбанк",
        "bank_account": "BG80UNCR12345678",
        "bank_branch_code": "UNCR",
        "bank_swift": "UNCRBGSF",
        "is_penalty": "False",
        "penalty_type": "",
        "penalty_amount": "",
        "penalty_percentage": "",
        "paid_through_date": "",
    },
    {
        "customer_egn": "9002025678",
        "loan_type_name": "Бизнес кредит",
        "app_date": "2024-01-10",
        "approval_date": "2024-01-12",
        "loan_amount": 50000.00,
        "term": 24,
        "installment_type": "monthly",
        "interest_rate": "",
        "start_date": "2024-02-01",
        "installment_start_date": "2024-02-01",
        "bank_name": "Пощенска банка",
        "bank_account": "BG90POST98765432",
        "bank_branch_code": "POST",
        "bank_swift": "BPBIBGSF",
        "is_penalty": "True",
        "penalty_type": "fixed",
        "penalty_amount": 50.00,
        "penalty_percentage": "",
        "paid_through_date": "2025-01-31",
    },
]


def generate_template(filename="loan_import_template.xlsx"):
    """Generate an empty Excel template with headers and 2 example rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loans"

    # Write header
    for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = openpyxl.styles.Font(bold=True)

    # Write example rows
    for row_idx, example in enumerate(EXAMPLE_ROWS, start=2):
        for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=example.get(col_name, ""))

    # Auto-adjust column widths
    for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(
            len(col_name) + 4, 18
        )

    wb.save(filename)
    print(f"Template saved to: {filename}")


# ──────────────────────────────────────────────
# VALIDATION HELPERS
# ──────────────────────────────────────────────

REQUIRED_FIELDS = [
    "customer_egn",
    "loan_type_name",
    "app_date",
    "approval_date",
    "loan_amount",
    "term",
    "installment_type",
    "start_date",
    "bank_name",
    "bank_account",
    "bank_branch_code",
    "bank_swift",
]


def validate_egn(egn):
    """Validate Bulgarian EGN format (10 digits with checksum)."""
    if not egn or not isinstance(egn, str):
        return False
    egn = egn.strip()
    if len(egn) != 10 or not egn.isdigit():
        return False
    # Checksum weights
    weights = [2, 4, 8, 5, 10, 9, 7, 3, 6]
    checksum = sum(int(egn[i]) * weights[i] for i in range(9)) % 11
    if checksum == 10:
        checksum = 0
    return checksum == int(egn[9])


def parse_date(value):
    """Parse a date value from Excel cell (may be string, datetime, or date)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
    return None


def parse_bool(value):
    """Parse a boolean value from Excel cell."""
    if value is None or value == "":
        return False
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("true", "1", "yes", "да")


def parse_float(value):
    """Parse a float value from Excel cell."""
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def parse_int(value):
    """Parse an integer value from Excel cell."""
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def validate_row(row_data, row_num):
    """Validate a single row. Returns (parsed_data, errors, warnings)."""
    errors = []
    warnings = []
    parsed = {}

    # Check required fields
    for field in REQUIRED_FIELDS:
        val = row_data.get(field)
        if val is None or (isinstance(val, str) and val.strip() == ""):
            errors.append(f"Missing required field: {field}")

    # Parse and validate EGN
    egn = str(row_data.get("customer_egn", "")).strip()
    # Handle numeric EGN that lost leading zero
    if egn and not egn.startswith("0") and len(egn) == 9:
        egn = "0" + egn
    parsed["customer_egn"] = egn
    if egn and not validate_egn(egn):
        warnings.append(f"EGN '{egn}' failed checksum validation (will still attempt lookup)")

    # Parse dates
    for date_field in ["app_date", "approval_date", "start_date", "installment_start_date",
                       "paid_through_date"]:
        parsed[date_field] = parse_date(row_data.get(date_field))

    if parsed["app_date"] and parsed["approval_date"]:
        if parsed["approval_date"] < parsed["app_date"]:
            errors.append("approval_date is before app_date")

    if parsed["app_date"] and parsed["start_date"]:
        if parsed["start_date"] < parsed["app_date"]:
            errors.append("start_date is before app_date")

    if not parsed["installment_start_date"]:
        parsed["installment_start_date"] = parsed.get("start_date")

    # Parse numeric fields
    parsed["loan_amount"] = parse_float(row_data.get("loan_amount"))
    if parsed["loan_amount"] is not None and parsed["loan_amount"] <= 0:
        errors.append("loan_amount must be greater than zero")

    parsed["term"] = parse_int(row_data.get("term"))
    if parsed["term"] is not None and parsed["term"] <= 0:
        errors.append("term must be greater than zero")

    parsed["interest_rate"] = parse_float(row_data.get("interest_rate"))

    # Installment type
    inst_type = str(row_data.get("installment_type", "monthly")).strip().lower()
    if inst_type not in ("monthly", "quarterly", "yearly"):
        errors.append(f"Invalid installment_type: '{inst_type}' (must be monthly/quarterly/yearly)")
    parsed["installment_type"] = inst_type

    # Loan type name
    parsed["loan_type_name"] = str(row_data.get("loan_type_name", "")).strip()

    # Bank details
    parsed["bank_name"] = str(row_data.get("bank_name", "")).strip()
    parsed["bank_account"] = str(row_data.get("bank_account", "")).strip()
    parsed["bank_branch_code"] = str(row_data.get("bank_branch_code", "")).strip()
    parsed["bank_swift"] = str(row_data.get("bank_swift", "")).strip()

    # Penalty
    parsed["is_penalty"] = parse_bool(row_data.get("is_penalty"))
    parsed["penalty_type"] = str(row_data.get("penalty_type", "fixed")).strip().lower() or "fixed"
    parsed["penalty_amount"] = parse_float(row_data.get("penalty_amount")) or 0.0
    parsed["penalty_percentage"] = parse_float(row_data.get("penalty_percentage")) or 0.0

    if parsed["is_penalty"]:
        if parsed["penalty_type"] not in ("fixed", "percentage"):
            errors.append(f"Invalid penalty_type: '{parsed['penalty_type']}'")
        if parsed["penalty_type"] == "fixed" and parsed["penalty_amount"] <= 0:
            warnings.append("Penalty is enabled with type 'fixed' but penalty_amount is 0")
        if parsed["penalty_type"] == "percentage" and parsed["penalty_percentage"] <= 0:
            warnings.append("Penalty is enabled with type 'percentage' but penalty_percentage is 0")

    return parsed, errors, warnings


# ──────────────────────────────────────────────
# ODOO XML-RPC HELPERS
# ──────────────────────────────────────────────

class OdooRPC:
    """Simple Odoo XML-RPC wrapper."""

    def __init__(self, url, db, user, password):
        self.url = url
        self.db = db
        self.common = xmlrpc.client.ServerProxy(f"{url}/xmlrpc/2/common")
        self.models = xmlrpc.client.ServerProxy(f"{url}/xmlrpc/2/object")
        self.uid = self.common.authenticate(db, user, password, {})
        if not self.uid:
            raise ConnectionError(f"Authentication failed for user '{user}' on database '{db}'")
        self.password = password
        print(f"Connected to {url} (db={db}) as uid={self.uid}")

    def execute(self, model, method, *args, **kwargs):
        """Execute an Odoo model method."""
        return self.models.execute_kw(
            self.db, self.uid, self.password, model, method, list(args), kwargs
        )

    def search(self, model, domain, **kwargs):
        return self.execute(model, "search", domain, **kwargs)

    def read(self, model, ids, fields=None):
        return self.execute(model, "read", ids, {"fields": fields} if fields else {})

    def search_read(self, model, domain, fields=None, **kwargs):
        kw = {}
        if fields:
            kw["fields"] = fields
        kw.update(kwargs)
        return self.execute(model, "search_read", domain, **kw)

    def create(self, model, vals):
        return self.execute(model, "create", [vals])

    def write(self, model, ids, vals):
        return self.execute(model, "write", ids, vals)

    def call(self, model, method, ids):
        """Call a button/action method on records."""
        return self.models.execute_kw(
            self.db, self.uid, self.password, model, method, [ids]
        )


# ──────────────────────────────────────────────
# LOOKUP FUNCTIONS
# ──────────────────────────────────────────────

def lookup_customer(rpc, egn):
    """Look up customer by personal_number (EGN). Returns (id, name) or (None, error)."""
    results = rpc.search_read(
        "res.partner",
        [("personal_number", "=", egn)],
        fields=["id", "name"],
        limit=1,
    )
    if not results:
        return None, f"No customer found with EGN '{egn}'"
    return results[0]["id"], results[0]["name"]


def lookup_loan_type(rpc, name):
    """Look up loan type by name. Returns (id, name) or (None, error)."""
    results = rpc.search_read(
        "customer.loan.type",
        [("name", "=", name)],
        fields=["id", "name"],
        limit=1,
    )
    if not results:
        return None, f"No loan type found with name '{name}'"
    return results[0]["id"], results[0]["name"]


def lookup_accounting_config(rpc):
    """Look up accounting accounts and journals needed for loan processing.

    Returns dict with keys: receivable_account_id, bank_cash_account,
    interest_income_account_id, journal_item_id, repayment_journal_item_id.
    Or (None, error).
    """
    config = {}
    errors = []

    # Receivable account
    accs = rpc.search_read(
        "account.account",
        [("account_type", "=", "asset_receivable")],
        fields=["id", "name"],
        limit=1,
    )
    if accs:
        config["receivable_account_id"] = accs[0]["id"]
    else:
        errors.append("No receivable account (asset_receivable) found")

    # Bank/Cash account
    accs = rpc.search_read(
        "account.account",
        [("account_type", "=", "asset_cash")],
        fields=["id", "name"],
        limit=1,
    )
    if accs:
        config["bank_cash_account"] = accs[0]["id"]
    else:
        errors.append("No bank/cash account (asset_cash) found")

    # Interest income account (also uses asset_receivable type in the model)
    accs = rpc.search_read(
        "account.account",
        [("account_type", "=", "asset_receivable")],
        fields=["id", "name"],
    )
    if len(accs) >= 2:
        # Use second receivable account for interest income if available
        config["interest_income_account_id"] = accs[1]["id"]
    elif accs:
        config["interest_income_account_id"] = accs[0]["id"]
    else:
        errors.append("No interest income account found")

    # Disbursement journal (type = 'general' or 'bank')
    journals = rpc.search_read(
        "account.journal",
        [("type", "in", ["general", "bank"])],
        fields=["id", "name", "type", "code"],
    )
    if journals:
        # Prefer bank journal for disbursement
        bank_journals = [j for j in journals if j["type"] == "bank"]
        config["journal_item_id"] = bank_journals[0]["id"] if bank_journals else journals[0]["id"]
        # Prefer general journal for repayment
        general_journals = [j for j in journals if j["type"] == "general"]
        config["repayment_journal_item_id"] = (
            general_journals[0]["id"] if general_journals else journals[0]["id"]
        )
    else:
        errors.append("No suitable journals found (need 'general' or 'bank' type)")

    if errors:
        return None, "; ".join(errors)
    return config, None


# ──────────────────────────────────────────────
# MAIN IMPORT LOGIC
# ──────────────────────────────────────────────

def read_excel(filepath):
    """Read the Excel file and return list of row dicts."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    # Read header row
    headers = []
    for cell in ws[1]:
        headers.append(str(cell.value).strip() if cell.value else "")

    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        # Skip entirely empty rows
        if all(v is None or v == "" for v in row):
            continue
        row_dict = {}
        for col_idx, value in enumerate(row):
            if col_idx < len(headers) and headers[col_idx]:
                row_dict[headers[col_idx]] = value
        row_dict["_row_num"] = row_idx
        rows.append(row_dict)

    return rows


def process_loan(rpc, parsed, row_num, accounting_config, test_mode=True):
    """Process a single loan row. Returns (success, loan_name_or_error)."""

    cutoff = parsed.get("paid_through_date")
    if cutoff is None:
        cutoff = date.today() if CUTOFF_DATE is None else datetime.strptime(CUTOFF_DATE, "%Y-%m-%d").date()

    # 1. Lookup customer
    customer_id, customer_info = lookup_customer(rpc, parsed["customer_egn"])
    if customer_id is None:
        return False, customer_info

    # 2. Lookup loan type
    loan_type_id, loan_type_info = lookup_loan_type(rpc, parsed["loan_type_name"])
    if loan_type_id is None:
        return False, loan_type_info

    if test_mode:
        # Count how many installments would be paid
        # We can't know exact count without creating, but estimate from term and dates
        inst_count = parsed["term"]
        start = parsed["installment_start_date"] or parsed["start_date"]
        paid_count = 0
        if start and cutoff:
            from dateutil.relativedelta import relativedelta
            current = start
            for i in range(inst_count):
                if current <= cutoff:
                    paid_count += 1
                if parsed["installment_type"] == "monthly":
                    current = current + relativedelta(months=1)
                elif parsed["installment_type"] == "quarterly":
                    current = current + relativedelta(months=3)
                else:
                    current = current + relativedelta(years=1)

        print(f"  [TEST] Would create loan:")
        print(f"    Customer: {customer_info} (id={customer_id})")
        print(f"    Loan Type: {loan_type_info} (id={loan_type_id})")
        print(f"    Amount: {parsed['loan_amount']}, Term: {parsed['term']} {parsed['installment_type']}")
        print(f"    Start Date: {parsed['start_date']}")
        print(f"    Interest Rate: {parsed['interest_rate'] or '(from loan type)'}")
        print(f"    Est. Installments: {inst_count}, Est. Paid (before {cutoff}): {paid_count}")
        print(f"    Bank: {parsed['bank_name']} / {parsed['bank_account']}")
        if parsed["is_penalty"]:
            print(f"    Penalty: {parsed['penalty_type']} "
                  f"({parsed['penalty_amount'] if parsed['penalty_type'] == 'fixed' else str(parsed['penalty_percentage']) + '%'})")
        return True, f"[TEST OK] Customer={customer_info}"

    # ── LIVE MODE ──

    # 3. Build loan vals
    loan_vals = {
        "customer_id": customer_id,
        "approved_loan_type_id": loan_type_id,
        "app_date": str(parsed["app_date"]),
        "approval_date": str(parsed["approval_date"]),
        "loan_amount": parsed["loan_amount"],
        "requested_loan_amount": parsed["loan_amount"],
        "term": parsed["term"],
        "requested_term": parsed["term"],
        "installment_type": parsed["installment_type"],
        "requested_installment_type": parsed["installment_type"],
        "start_date": str(parsed["start_date"]),
        "requested_start_date": str(parsed["start_date"]),
        "installment_start_date": str(parsed["installment_start_date"] or parsed["start_date"]),
        "cst_bank_name": parsed["bank_name"],
        "cst_bank_account_number": parsed["bank_account"],
        "cst_bank_branch_code": parsed["bank_branch_code"],
        "cst_bank_swift_bic_code": parsed["bank_swift"],
        "is_penalty": parsed["is_penalty"],
        "status": "draft",
        # Accounting fields
        "receivable_account_id": accounting_config["receivable_account_id"],
        "bank_cash_account": accounting_config["bank_cash_account"],
        "interest_income_account_id": accounting_config["interest_income_account_id"],
        "journal_item_id": accounting_config["journal_item_id"],
        "repayment_journal_item_id": accounting_config["repayment_journal_item_id"],
    }

    if parsed["interest_rate"] is not None:
        loan_vals["interest_rate"] = parsed["interest_rate"]

    if parsed["is_penalty"]:
        loan_vals["penalty_type"] = parsed["penalty_type"]
        loan_vals["penalty_amount"] = parsed["penalty_amount"]
        loan_vals["penalty_percentage"] = parsed["penalty_percentage"]

    # 4. Create loan record
    loan_id = rpc.create("customer.loan", loan_vals)
    print(f"  Created loan id={loan_id}")

    # Read back the generated name
    loan_data = rpc.read("customer.loan", [loan_id], ["name"])[0]
    loan_name = loan_data["name"]
    print(f"  Loan name: {loan_name}")

    # 5. Compute installment schedule
    rpc.call("customer.loan", "compute_installment", [loan_id])
    print(f"  Computed installment schedule")

    # 6. Progress to disbursement status
    rpc.write("customer.loan", [loan_id], {"status": "disbursement"})

    # 7. Call action_disburse_loan to create disbursement journal entry
    rpc.call("customer.loan", "action_disburse_loan", [loan_id])
    print(f"  Created disbursement journal entry")

    # 8. Post the disbursement journal entry
    loan_data = rpc.read("customer.loan", [loan_id], ["journal_entry_id"])[0]
    je_id = loan_data["journal_entry_id"]
    if je_id:
        je_id = je_id[0] if isinstance(je_id, (list, tuple)) else je_id
        rpc.call("account.move", "action_post", [je_id])
        print(f"  Posted disbursement journal entry (id={je_id})")

    # 9. Set status to in_progress
    rpc.write("customer.loan", [loan_id], {"status": "in_progress"})
    print(f"  Status set to 'in_progress'")

    # 10. Process past installments
    loan_lines = rpc.search_read(
        "customer.loan.lines",
        [("customer_loan_id", "=", loan_id), ("display_type", "=", False)],
        fields=["id", "emi_date", "installments_no"],
        order="emi_date asc",
    )

    paid_count = 0
    for line in loan_lines:
        emi_date = line["emi_date"]
        if isinstance(emi_date, str):
            emi_date = datetime.strptime(emi_date, "%Y-%m-%d").date()

        if emi_date <= cutoff:
            line_id = line["id"]

            # Create journal entry for this installment
            rpc.call("customer.loan.lines", "action_create_journal_entry", [line_id])

            # Read back the journal entry id
            line_data = rpc.read("customer.loan.lines", [line_id], ["journal_entry_id"])[0]
            line_je_id = line_data.get("journal_entry_id")
            if line_je_id:
                line_je_id = line_je_id[0] if isinstance(line_je_id, (list, tuple)) else line_je_id
                # Set journal entry date to emi_date
                rpc.write("account.move", [line_je_id], {"date": str(emi_date)})
                # Post the journal entry
                rpc.call("account.move", "action_post", [line_je_id])

            paid_count += 1

    print(f"  Paid {paid_count}/{len(loan_lines)} installments (cutoff={cutoff})")
    return True, loan_name


def main():
    if "--template" in sys.argv:
        generate_template()
        return

    print(f"{'=' * 60}")
    print(f"Loan Import Script {'(TEST MODE)' if TEST_MODE else '(LIVE MODE)'}")
    print(f"{'=' * 60}")
    print(f"Excel file: {EXCEL_FILE}")
    print()

    # Read Excel
    try:
        rows = read_excel(EXCEL_FILE)
    except FileNotFoundError:
        print(f"ERROR: File not found: {EXCEL_FILE}")
        print("Run with --template to generate a sample Excel file.")
        sys.exit(1)
    except Exception as e:
        print(f"ERROR reading Excel: {e}")
        sys.exit(1)

    if not rows:
        print("No data rows found in Excel file.")
        sys.exit(0)

    print(f"Found {len(rows)} data row(s)")
    print()

    # Validate all rows first
    all_parsed = []
    validation_errors = 0

    for row_data in rows:
        row_num = row_data.get("_row_num", "?")
        parsed, errors, warnings = validate_row(row_data, row_num)

        if warnings:
            for w in warnings:
                print(f"  [WARN] Row {row_num}: {w}")

        if errors:
            validation_errors += 1
            for e in errors:
                print(f"  [ERROR] Row {row_num}: {e}")
            all_parsed.append((row_num, None, errors))
        else:
            all_parsed.append((row_num, parsed, []))

    print()
    if validation_errors:
        print(f"Validation: {validation_errors} row(s) have errors")
    print()

    # Connect to Odoo
    try:
        rpc = OdooRPC(ODOO_URL, ODOO_DB, ODOO_USER, ODOO_PASS)
    except Exception as e:
        print(f"ERROR connecting to Odoo: {e}")
        sys.exit(1)
    print()

    # Lookup accounting config once
    accounting_config, acct_error = lookup_accounting_config(rpc)
    if acct_error:
        print(f"ERROR: Accounting config issue: {acct_error}")
        if not TEST_MODE:
            print("Cannot proceed in live mode without accounting configuration.")
            sys.exit(1)
        else:
            print("(Continuing in test mode despite accounting config issues)")
            accounting_config = accounting_config or {}
    print()

    # Process each row
    created = 0
    failed = 0
    skipped = 0

    for row_num, parsed, errors in all_parsed:
        print(f"Row {row_num}:")
        if parsed is None:
            print(f"  SKIPPED (validation errors: {'; '.join(errors)})")
            skipped += 1
            continue

        try:
            success, result = process_loan(rpc, parsed, row_num, accounting_config, TEST_MODE)
            if success:
                print(f"  {'[TEST] OK' if TEST_MODE else 'SUCCESS'}: {result}")
                created += 1
            else:
                print(f"  FAILED: {result}")
                failed += 1
        except Exception as e:
            print(f"  FAILED (exception): {e}")
            failed += 1

        print()

    # Summary
    print(f"{'=' * 60}")
    print(f"SUMMARY {'(TEST MODE)' if TEST_MODE else '(LIVE MODE)'}:")
    print(f"  {'Would create' if TEST_MODE else 'Created'}: {created}")
    print(f"  Failed: {failed}")
    print(f"  Skipped (validation): {skipped}")
    print(f"  Total rows: {len(rows)}")
    print(f"{'=' * 60}")


if __name__ == "__main__":
    main()
